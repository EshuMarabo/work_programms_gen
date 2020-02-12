import docx
import settings
import re;
import data_classes  as dc
from docx.table import Table
from docx.text.paragraph import Paragraph 
import os.path




def old_work_program_parse(path=None):
    try:
        
        if path==None:
            path=settings.cfg.temp_path
        if (not os.path.exists(path)):
            return None;
        doc = docx.Document(path);
        elements=[]
        for e in doc.element.iterchildren():
            for element in e.iterchildren():
                elements.append((element,e))
                
        indexes=range(0,len(elements)-1)
        for_return = dc.course_content();
        for i in indexes:
            if type(elements[i][0]) == docx.oxml.text.paragraph.CT_P:
                paragr = Paragraph(elements[i][0],elements[i][1])
        
                if re.match("4.2. Н",paragr.text) is not None:
                    #print(paragr.text)
                    if type(elements[i+1][0]) == docx.oxml.table.CT_Tbl:
                        table = Table(elements[i+1][0],elements[i+1][1])
                        row_number=1;
                        while row_number<len(table.rows):
                            theme_name = table.cell(row_number,1).text
                            theme_content = table.cell(row_number,2).text
                            theme_number = table.cell(row_number,0).text
                            for_return.themes.append(dc.cource_part(theme_name,theme_content,theme_number))
                            row_number+=1
                them_indexer =range(0,len(for_return.themes))
        
                if re.match("4.3. Л",paragr.text) is not None:
                    if type(elements[i+1][0]) == docx.oxml.table.CT_Tbl:
                        table = Table(elements[i+1][0],elements[i+1][1])
                        row_number=1;
                        while row_number<len(table.rows):
                            lab_theme_content = table.cell(row_number,2).text
                            lab_theme_number = table.cell(row_number,1).text
                            for ti in them_indexer:
                                if lab_theme_number!='':
                                    if for_return.themes[ti].number== int(lab_theme_number):
                                        for_return.labs.append(dc.lab_work(for_return.themes[ti],lab_theme_content))
                            row_number+=1       
        
                if re.match("4.4. П",paragr.text) is not None:
                    if type(elements[i+1][0]) == docx.oxml.table.CT_Tbl:
                        table = Table(elements[i+1][0],elements[i+1][1])
                        row_number=1;
                        while row_number<len(table.rows):
                            pract_theme_content = table.cell(row_number,2).text
                            pract_theme_number = table.cell(row_number,1).text
                            for ti in them_indexer:
                                if pract_theme_number!='':
                                    if for_return.themes[ti].number== int(pract_theme_number):
                                        for_return.practic.append(dc.practical(for_return.themes[ti],pract_theme_content))
                            row_number+=1
                
        
                if paragr.text!='':
                    if re.match("а\) О",paragr.text) is not None:
                        row_number=1;
                        while (row_number<len(elements)):
                            par=Paragraph(elements[i+row_number][0],elements[i+row_number][1])
                            #print(par.text)
                            if re.match("б\) Д",par.text) is not None:
                                break;
                            if par.text=='':
                                break;
                            for_return. main_lit.append(dc.disciplin._edit_source(par.text));
                            row_number+=1;
                
        
                if paragr.text!='':
                    if re.match("б\) Д",paragr.text) is not None:
                        row_number=1;
                        while (row_number<len(elements)) and (
                                Paragraph(elements[i+row_number][0],elements[i+row_number][1]).text!=""):
                            par=Paragraph(elements[i+row_number][0],elements[i+row_number][1])
                            if par.text=='':
                                break;
                            for_return.additional_lit.append(dc.disciplin._edit_source(par.text));
                            row_number+=1;
                if paragr.text!='':
                    if re.match("Опрос проводится в устной ",paragr.text) is not None:
        
                        row_number=1;
                        while (row_number<len(elements)) and (
                                Paragraph(elements[i+row_number][0],elements[i+row_number][1]).text!=""):
                            par=Paragraph(elements[i+row_number][0],elements[i+row_number][1])
                            if par.text=='':
                                break;
                            for_return.current_questions.append(dc.disciplin._edit_source(par.text));
                            row_number+=1;     
                            
                if paragr.text!='':
                    if re.match("Перечень вопросов для подготовки к экзамену:",paragr.text) is not None:
        
                        row_number=1;
                        while (row_number<len(elements)) and (
                                Paragraph(elements[i+row_number][0],elements[i+row_number][1]).text!=""):
                            par=Paragraph(elements[i+row_number][0],elements[i+row_number][1])
                            if par.text=='':
                                break;
                            for_return.control_questions.append(dc.disciplin._edit_source(par.text));
                            row_number+=1;                   
        return for_return;
    except:
        return None;
           
                    
               # print(str(len(table.rows))+", "+str(len(table.columns)))
                #print(table.cell(0,0).text)
    
        
                
            
    
